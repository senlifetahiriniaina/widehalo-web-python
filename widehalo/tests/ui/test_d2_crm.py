"""D-2 (CRM) — ce que l'ecran des opportunites AFFICHE, et CRM-5 au complet.

Deux ecrans du meme objet disaient deux choses : la liste rendait
`Pipeline commercial par defaut:appointment_scheduled`, la fiche « Rendez-vous
planifié ». Et CRM-5 etait compte tenu sur sa moitie visible — l'etat vide
proposait la creation, jamais l'import, parce qu'aucune route d'import
n'existait vers laquelle pointer.
"""

from __future__ import annotations

import io

import pytest
from apps.core.models.tenant import Tenant
from apps.core.models.user import User
from apps.core.tests.utils import grant_role, use_tenant
from apps.crm.models import CrmLead, CrmStage
from apps.crm.services.leads import create_lead_quick
from apps.crm.services.pipelines import ensure_default_pipeline
from apps.partners.services.onboarding import create_partner
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import Client
from django.test.utils import CaptureQueriesContext

pytestmark = pytest.mark.django_db


def _classeur(lignes: list[list[object]], entete: list[str]) -> bytes:
    from openpyxl import Workbook

    classeur = Workbook()
    feuille = classeur.active
    feuille.append(entete)
    for ligne in lignes:
        feuille.append(ligne)
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


def _televerse(octets: bytes) -> SimpleUploadedFile:
    return SimpleUploadedFile(
        "opportunites.xlsx",
        octets,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


ENTETE = ["Nom", "Tiers", "Contact", "Email", "Téléphone", "Source", "Montant attendu (MGA)"]


@pytest.fixture
def crm():
    tenant = Tenant.objects.create(code="D2-CRM", name="D2 CRM Tenant")
    with use_tenant(tenant.id):
        user = User.objects.create_user(email="d2-crm@example.com", password="Str0ngPassw0rd!23")
        grant_role(user, "commercial")
        ensure_default_pipeline(tenant)
        lead = create_lead_quick(tenant=tenant, name="Affaire D2", salesperson=user)
    client = Client()
    client.force_login(user)
    session = client.session
    session["tenant_id"] = str(tenant.id)
    session.save()
    return client, tenant, user, lead


def test_la_liste_nomme_l_etape_comme_la_fiche(crm) -> None:
    client, tenant, _user, lead = crm
    with use_tenant(tenant.id):
        nom_etape = lead.stage.name

    liste = client.get("/crm/").content.decode()
    fiche = client.get(f"/crm/{lead.id}/").content.decode()

    assert nom_etape in liste
    assert nom_etape in fiche
    # Le `__str__` technique de `CrmStage` ne doit plus atteindre l'ecran.
    assert f":{lead.stage.code}" not in liste


def test_l_etape_ne_coute_pas_une_requete_par_ligne(crm) -> None:
    """`select_related` n'est pas du confort : sans lui, chaque ligne relit
    son etape — et l'ancienne colonne relisait EN PLUS le pipeline, par le
    `__str__` de `CrmStage`."""
    client, tenant, user, _lead = crm
    with use_tenant(tenant.id):
        for numero in range(2, 12):
            create_lead_quick(tenant=tenant, name=f"Affaire D2 n{numero}", salesperson=user)

    with CaptureQueriesContext(connection) as requetes:
        assert client.get("/crm/").status_code == 200
    lectures_etape = [
        requete["sql"]
        for requete in requetes.captured_queries
        if "crm_stage" in requete["sql"] and "COUNT" not in requete["sql"].upper()
    ]
    assert len(lectures_etape) <= 2, (
        f"{len(lectures_etape)} lectures de crm_stage pour 11 lignes :\n"
        + "\n".join(lectures_etape[:5])
    )


def test_l_etat_vide_propose_les_deux_voies_du_critere(crm) -> None:
    """CRM-5, mot pour mot : « proposant la creation d'une opportunite ET
    l'import d'un fichier ». Le composant portait `import_url` depuis L4 ;
    il n'y avait aucune route vers laquelle pointer."""
    client, tenant, _user, lead = crm
    with use_tenant(tenant.id):
        CrmLead.objects.filter(id=lead.id).update(is_active=False)

    contenu = client.get("/crm/").content.decode()
    assert "/crm/new/" in contenu
    assert "/crm/imports/" in contenu


def test_un_import_cree_les_opportunites_et_rattache_les_tiers_connus(crm) -> None:
    client, tenant, _user, _lead = crm
    with use_tenant(tenant.id):
        tiers = create_partner(tenant=tenant, name="Client D2 SARL", roles=["client"])

    fichier = _classeur(
        [
            ["Extension entrepot", "Client D2 SARL", "Rakoto", "a@b.mg", "+261", "Salon", 1500000],
            ["Renouvellement", "Societe Inconnue", "", "", "", "", 250000],
        ],
        ENTETE,
    )
    reponse = client.post("/crm/imports/", {"file": _televerse(fichier)})
    assert reponse.status_code == 200

    with use_tenant(tenant.id):
        creees = {lead.name: lead for lead in CrmLead.objects.filter(is_active=True)}
    assert "Extension entrepot" in creees
    assert "Renouvellement" in creees
    assert creees["Extension entrepot"].partner_id == tiers.id
    # Un nom inconnu n'est pas une erreur : l'opportunite existe, sans tiers.
    assert creees["Renouvellement"].partner_id is None
    assert "1 nom(s) de tiers sans correspondance" in reponse.content.decode().replace(
        "&#x27;", "'"
    )


def test_un_fichier_fautif_n_ecrit_rien_et_nomme_la_ligne_du_tableur(crm) -> None:
    """Lecon BNK-2, payee une fois : un import qui ecrit les lignes valides
    puis leve laisse un lot a moitie charge, un refus, et des doublons au
    re-essai. Et le numero rendu est celui que l'utilisateur LIT dans son
    tableur — l'import des partenaires rend « Ligne 0 » pour sa premiere
    ligne de donnees."""
    client, tenant, _user, _lead = crm
    # **Hors contexte de societe, `objects` ne rend RIEN** : compter ici
    # donnait zero avant comme apres, et l'assertion etait vraie quoi qu'il
    # arrive. C'est la falsification F131 qui l'a demasque en NE MORDANT
    # PAS — une falsification qui ne mord pas doit etre crue.
    with use_tenant(tenant.id):
        avant = CrmLead.objects.filter(is_active=True).count()

    fichier = _classeur(
        [
            ["Bonne ligne", "", "", "", "", "", 1000],
            ["", "", "", "", "", "", 2000],
            ["Montant illisible", "", "", "", "", "", "beaucoup"],
        ],
        ENTETE,
    )
    reponse = client.post("/crm/imports/", {"file": _televerse(fichier)})
    contenu = reponse.content.decode()

    with use_tenant(tenant.id):
        assert CrmLead.objects.filter(is_active=True).count() == avant
    assert "Ligne 3" in contenu, "la ligne vide est la 3e du tableur (1 = en-tete)"
    assert "Ligne 4" in contenu
    assert "Ligne 2" not in contenu, "la ligne valide n'a pas a etre signalee"


def test_l_import_ne_relit_pas_le_referentiel_ligne_par_ligne(crm) -> None:
    """`find_partner_by_name` charge TOUT le referentiel a chaque appel : un
    appel par ligne rendrait l'import quadratique. Le defaut a deja ete paye
    deux fois dans ce depot (T3, BNK-3) et se lit dans la fonction appelee,
    jamais dans son nom."""
    from apps.crm.services.lead_import import import_leads_xlsx

    _client, tenant, _user, _lead = crm
    with use_tenant(tenant.id):
        for numero in range(5):
            create_partner(tenant=tenant, name=f"Tiers D2 n{numero}", roles=["client"])
        fichier = _classeur(
            [[f"Affaire {n}", f"Tiers D2 n{n % 5}", "", "", "", "", 1000] for n in range(12)],
            ENTETE,
        )
        with CaptureQueriesContext(connection) as requetes:
            resume = import_leads_xlsx(tenant, fichier)

    assert resume.created_count == 12
    lectures = [
        requete["sql"]
        for requete in requetes.captured_queries
        if "partners_partner" in requete["sql"]
        and requete["sql"].lstrip().upper().startswith("SELECT")
    ]
    assert len(lectures) == 1, f"{len(lectures)} balayages du referentiel tiers"


def test_le_modele_telecharge_porte_les_memes_colonnes_que_le_lecteur(crm) -> None:
    """Deux listes d'en-tetes divergeraient a la premiere colonne ajoutee."""
    from apps.crm.services.lead_import import COLONNES
    from openpyxl import load_workbook

    client, *_ = crm
    reponse = client.get("/crm/imports/template/")
    assert reponse.status_code == 200
    feuille = load_workbook(io.BytesIO(reponse.content)).active
    entete = [cellule.value for cellule in next(feuille.iter_rows())]
    assert entete == [libelle for libelle, _champ in COLONNES]


def test_un_role_sans_droit_de_creation_n_ouvre_pas_l_ecran_d_import(crm) -> None:
    """C-1d : un ecran dont la seule raison d'etre est la creation exige le
    droit de creer pour S'OUVRIR, pas seulement pour se soumettre."""
    _client, tenant, _user, _lead = crm
    with use_tenant(tenant.id):
        lecteur = User.objects.create_user(
            email="d2-lecture@example.com", password="Str0ngPassw0rd!23"
        )
        grant_role(lecteur, "controleur_gestion")
    client = Client()
    client.force_login(lecteur)
    session = client.session
    session["tenant_id"] = str(tenant.id)
    session.save()

    assert client.get("/crm/imports/").status_code == 403


def test_les_etapes_semees_sont_accentuees(crm) -> None:
    _client, tenant, _user, _lead = crm
    with use_tenant(tenant.id):
        noms = set(CrmStage.objects.values_list("name", flat=True))
    assert "Rendez-vous planifié" in noms
    assert "Rendez-vous planifie" not in noms
    assert "Gagné" in noms
