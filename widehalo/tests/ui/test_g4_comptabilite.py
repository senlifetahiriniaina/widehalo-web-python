"""G-4 — l'import de factures, les coûts d'approche, la DCOM, l'analytique
et les taux de change : cinq chaînes qui n'avaient pas d'écran.

**Et un défaut que la mesure a trouvé en écrivant ce lot.**
`record_analytic_lines` n'avait **aucun appelant de production** :
`enforce_and_validate` validait la distribution à la saisie, et rien ne la
matérialisait. `AccAnalyticLine` restait vide, et les deux lectures qui s'y
appuient — le compte de résultat analytique et l'écart budgétaire par axe —
rendaient zéro partout, sans la moindre erreur. Un rapport qui rend zéro
ressemble à une société sans activité sur cet axe : c'est la variante la
plus discrète du motif « rien de décoratif ».
"""

from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

import pytest
from apps.accounting.models import (
    AccAccount,
    AccAnalyticAccount,
    AccAnalyticLine,
    AccAnalyticPlan,
    AccDcomDeclaration,
    AccExchangeRate,
    AccFiscalYear,
    AccInvoiceImportBatch,
    AccJournal,
    AccLandedCostBatch,
    AccPeriod,
)
from apps.accounting.services.currency import get_rate
from apps.accounting.services.moves import add_line, create_draft_move, post_move
from apps.core.models.tenant import Tenant
from apps.core.models.user import User, UserTenantMembership
from apps.core.tests.utils import grant_module_access, use_tenant
from apps.partners.models import Partner
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from openpyxl import Workbook

pytestmark = pytest.mark.django_db

_ENTETE = [
    "REFERENCE",
    "DATE",
    "SENS",
    "PARTENAIRE",
    "CODE_PRODUIT",
    "DESIGNATION",
    "QUANTITE",
    "PRIX_UNITAIRE",
    "TAUX_TVA",
    "COMPTE",
]


def _xlsx(lignes: list[list[object]]) -> bytes:
    classeur = Workbook()
    feuille = classeur.active
    feuille.append(_ENTETE)
    for ligne in lignes:
        feuille.append(ligne)
    tampon = io.BytesIO()
    classeur.save(tampon)
    return tampon.getvalue()


@pytest.fixture
def societe():
    tenant = Tenant.objects.create(code="G4", name="Societe couts")
    user = User.objects.create_user(email="g4@example.com", password="Str0ngPassw0rd!23")
    grant_module_access(user, "accounting")
    # `qualify_accinvoiceimportrow` est une permission CUSTOM : elle ne
    # commence par aucun des trois prefixes que `grant_module_access`
    # accorde, et sans elle l'ecran de qualification refuserait.
    from django.contrib.auth.models import Permission

    user.user_permissions.add(
        *Permission.objects.filter(
            codename="qualify_accinvoiceimportrow", content_type__app_label="accounting"
        )
    )
    UserTenantMembership.objects.get_or_create(
        user=user, tenant=tenant, defaults={"is_default": True}
    )
    with use_tenant(tenant.id):
        exercice = AccFiscalYear.objects.create(
            tenant=tenant,
            code="2026",
            date_start=dt.date(2026, 1, 1),
            date_end=dt.date(2026, 12, 31),
        )
        periode = AccPeriod.objects.create(
            tenant=tenant,
            fiscal_year=exercice,
            code="2026-01",
            date_start=dt.date(2026, 1, 1),
            date_end=dt.date(2026, 1, 31),
        )
        journal = AccJournal.objects.create(
            tenant=tenant,
            code="OD",
            name="Operations diverses",
            type=AccJournal.TYPE_MISC,
            sequence_prefix="OD",
        )
        AccJournal.objects.create(
            tenant=tenant,
            code="ACH",
            name="Achats",
            type=AccJournal.TYPE_PURCHASE,
            sequence_prefix="ACH",
        )
        AccAccount.objects.create(
            tenant=tenant,
            code="401000",
            name="Fournisseurs",
            account_class="4",
            type=AccAccount.TYPE_PAYABLE,
        )
        charge = AccAccount.objects.create(
            tenant=tenant,
            code="607000",
            name="Achats de marchandises",
            account_class="6",
            type=AccAccount.TYPE_EXPENSE,
        )
        produit = AccAccount.objects.create(
            tenant=tenant,
            code="707000",
            name="Ventes de marchandises",
            account_class="7",
            type=AccAccount.TYPE_INCOME,
        )
        fournisseur = Partner.objects.create(
            tenant=tenant,
            name="Fournisseur Andry",
            nif="9876543210",
            roles=[Partner.ROLE_SUPPLIER],
        )
    client = Client()
    client.force_login(user)
    session = client.session
    session["tenant_id"] = str(tenant.id)
    session.save()
    return {
        "tenant": tenant,
        "user": user,
        "client": client,
        "exercice": exercice,
        "periode": periode,
        "journal": journal,
        "charge": charge,
        "produit": produit,
        "fournisseur": fournisseur,
    }


# --------------------------------------------------------------------------
# Import de factures fournisseur
# --------------------------------------------------------------------------


def test_un_fichier_de_factures_se_depose_et_son_rapport_nomme_chaque_ligne(societe) -> None:
    fichier = _xlsx(
        [
            [
                "FAC-001",
                dt.date(2026, 1, 15),
                "fournisseur",
                "Fournisseur Andry",
                "PRD-INCONNU",
                "Tissu coton",
                10,
                15000,
                20,
                "607000",
            ]
        ]
    )
    reponse = societe["client"].post(
        "/accounting/config/imports/invoices/",
        {"file": SimpleUploadedFile("factures.xlsx", fichier)},
    )
    assert reponse.status_code == 200
    contenu = reponse.content.decode()
    assert "Résultat" in contenu

    with use_tenant(societe["tenant"].id):
        lot = AccInvoiceImportBatch.objects.first()
        assert lot is not None, "Le dépôt n'a créé aucun lot."
        assert lot.total_rows == 1
        ligne = lot.rows.first()
        assert ligne is not None
        assert ligne.invoice_reference == "FAC-001"

    detail = societe["client"].get(f"/accounting/config/imports/invoices/{lot.id}/")
    assert detail.status_code == 200
    assert "FAC-001" in detail.content.decode()


def test_le_modele_de_fichier_porte_les_entetes_que_le_service_lit(societe) -> None:
    """Un modèle qui diverge du format lu ne se voit qu'au premier import
    raté : le test compare le fichier proposé aux alias réellement acceptés."""
    from apps.accounting.services.invoice_import import INVOICE_IMPORT_HEADER_ALIASES

    reponse = societe["client"].get("/accounting/config/imports/invoices/template.xlsx")
    assert reponse.status_code == 200
    classeur = __import__("openpyxl").load_workbook(io.BytesIO(reponse.content))
    entetes = [cellule.value for cellule in next(classeur.active.iter_rows(max_row=1))]
    for champ, alias in INVOICE_IMPORT_HEADER_ALIASES.items():
        assert any(entete in alias for entete in entetes), (
            f"Le modèle proposé ne porte aucun en-tête reconnu pour le champ « {champ} »."
        )


# --------------------------------------------------------------------------
# Coûts d'approche
# --------------------------------------------------------------------------


def test_un_lot_de_couts_dapproche_repartit_ses_frais_avant_toute_finalisation(societe) -> None:
    reponse = societe["client"].post(
        "/accounting/landed-costs/",
        {
            "label": "Importation janvier",
            "date": "2026-01-20",
            "allocation_method": AccLandedCostBatch.METHOD_BY_VALUE,
        },
    )
    assert reponse.status_code == 302, reponse.content
    with use_tenant(societe["tenant"].id):
        lot = AccLandedCostBatch.objects.first()
    assert lot is not None

    societe["client"].post(
        f"/accounting/landed-costs/{lot.id}/",
        {
            "action": "add_line",
            "description": "Tissu coton",
            "qty": "100",
            "purchase_value_mga": "1000000",
        },
    )
    societe["client"].post(
        f"/accounting/landed-costs/{lot.id}/",
        {
            "action": "add_component",
            "component_label": "Fret maritime",
            "amount_mga": "200000",
        },
    )

    contenu = societe["client"].get(f"/accounting/landed-costs/{lot.id}/").content.decode()
    assert "Tissu coton" in contenu
    assert "Fret maritime" in contenu
    # Tous les frais tombent sur l'unique ligne : coût débarqué 1 200 000 Ar,
    # soit 12 000 Ar l'unité pour 100 unités.
    assert "1 200 000 Ar" in contenu
    assert "12 000 Ar" in contenu

    with use_tenant(societe["tenant"].id):
        lot.refresh_from_db()
        assert lot.state == AccLandedCostBatch.STATE_DRAFT, (
            "Le lot s'est finalisé sans qu'on le demande."
        )

    finalisation = societe["client"].post(
        f"/accounting/landed-costs/{lot.id}/", {"action": "finalize"}
    )
    assert finalisation.status_code == 302
    with use_tenant(societe["tenant"].id):
        lot.refresh_from_db()
        assert lot.state == AccLandedCostBatch.STATE_FINALIZED

    ferme = societe["client"].post(
        f"/accounting/landed-costs/{lot.id}/",
        {
            "action": "add_line",
            "description": "Ajout tardif",
            "qty": "1",
            "purchase_value_mga": "1",
        },
    )
    assert ferme.status_code == 200
    assert "déjà finalisé" in ferme.content.decode()


# --------------------------------------------------------------------------
# DCOM
# --------------------------------------------------------------------------


def test_la_declaration_dcom_se_genere_et_nomme_ses_tiers(societe) -> None:
    with use_tenant(societe["tenant"].id):
        brouillon = create_draft_move(
            tenant=societe["tenant"],
            journal=societe["journal"],
            period=societe["periode"],
            date=dt.date(2026, 1, 10),
            narration="Achat",
        )
        ligne = add_line(
            brouillon, account=societe["charge"], label="Achat", debit=Decimal("300000")
        )
        add_line(brouillon, account=societe["produit"], label="Vente", credit=Decimal("300000"))
        ligne.partner_id = societe["fournisseur"].id
        ligne.save(update_fields=["partner_id"])
        post_move(brouillon)

    reponse = societe["client"].post(
        "/accounting/dcom/", {"fiscal_year_id": str(societe["exercice"].id)}
    )
    assert reponse.status_code == 302, reponse.content

    with use_tenant(societe["tenant"].id):
        declaration = AccDcomDeclaration.objects.first()
        assert declaration is not None
        assert declaration.total_amount_mga == Decimal("300000.0000")

    contenu = (
        societe["client"]
        .get(f"/accounting/dcom/?fiscal_year_id={societe['exercice'].id}")
        .content.decode()
    )
    assert "Fournisseur Andry" in contenu
    assert str(societe["fournisseur"].id) not in contenu


# --------------------------------------------------------------------------
# Analytique
# --------------------------------------------------------------------------


def test_un_axe_analytique_se_declare_puis_se_materialise_a_la_publication(societe) -> None:
    """La propriété qui ferme le défaut mesuré dans ce lot.

    Avant G-4 rien n'appelait `record_analytic_lines` : la distribution
    était validée puis oubliée, et `AccAnalyticLine` restait vide."""
    reponse = societe["client"].post(
        "/accounting/config/analytic-plans/", {"code": "projet", "name": "Projets"}
    )
    assert reponse.status_code == 200, reponse.content
    with use_tenant(societe["tenant"].id):
        plan = AccAnalyticPlan.objects.filter(code="projet").first()
    assert plan is not None

    reponse = societe["client"].post(
        "/accounting/config/analytic-accounts/",
        {"plan_id": str(plan.id), "code": "P-042", "name": "Chantier Nord"},
    )
    assert reponse.status_code == 200, reponse.content
    with use_tenant(societe["tenant"].id):
        assert AccAnalyticAccount.objects.filter(code="P-042").exists()

        brouillon = create_draft_move(
            tenant=societe["tenant"],
            journal=societe["journal"],
            period=societe["periode"],
            date=dt.date(2026, 1, 10),
            narration="Achat ventilé",
        )
        add_line(
            brouillon,
            account=societe["charge"],
            label="Achat",
            debit=Decimal("500000"),
            analytic_distribution={"projet": {"P-042": 100}},
        )
        add_line(brouillon, account=societe["produit"], label="Vente", credit=Decimal("500000"))
        post_move(brouillon)

        lignes = list(AccAnalyticLine.objects.all())
        assert len(lignes) == 1, (
            "La distribution analytique n'a pas été matérialisée à la publication."
        )
        assert lignes[0].amount == Decimal("500000.0000")


def test_la_materialisation_analytique_est_une_projection_pas_un_journal(societe) -> None:
    """La propriété que la passe complète a imposée.

    En branchant `record_analytic_lines` sur `post_move`, deux tests de la
    phase 2 ont **doublé** leurs montants : ils l'appelaient eux-mêmes
    après publication, et chaque appel ajoutait un jeu de lignes. Le défaut
    n'était pas dans les tests. Les lignes analytiques d'une ligne
    d'écriture sont la PROJECTION de sa distribution : deux projections de
    la même distribution donnent le même résultat, jamais le double."""
    from apps.accounting.services.analytics import record_analytic_lines

    with use_tenant(societe["tenant"].id):
        plan = AccAnalyticPlan.objects.create(
            tenant=societe["tenant"], code="atelier", name="Ateliers"
        )
        AccAnalyticAccount.objects.create(
            tenant=societe["tenant"], plan=plan, code="AT-1", name="Atelier Nord"
        )
        brouillon = create_draft_move(
            tenant=societe["tenant"],
            journal=societe["journal"],
            period=societe["periode"],
            date=dt.date(2026, 1, 10),
            narration="Achat ventilé",
        )
        ligne = add_line(
            brouillon,
            account=societe["charge"],
            label="Achat",
            debit=Decimal("400000"),
            analytic_distribution={"atelier": {"AT-1": 100}},
        )
        add_line(brouillon, account=societe["produit"], label="Vente", credit=Decimal("400000"))
        post_move(brouillon)

        assert AccAnalyticLine.objects.filter(move_line=ligne).count() == 1

        # Une seconde projection REMPLACE la première : elle ne s'y ajoute
        # pas. Sans cela, la garantie reposerait sur le fait qu'un seul
        # appelant existe — vrai hier, faux demain.
        record_analytic_lines(ligne)
        lignes = list(AccAnalyticLine.objects.filter(move_line=ligne))
        assert len(lignes) == 1
        assert lignes[0].amount == Decimal("400000.0000")


def test_un_axe_inconnu_refuse_la_publication_en_le_nommant(societe) -> None:
    """`AccAnalyticAccount.objects.get` levait `DoesNotExist` : un 500 à la
    publication d'une écriture dont la distribution désigne un code non
    déclaré. Le refus nomme désormais le plan et le compte."""
    with use_tenant(societe["tenant"].id):
        brouillon = create_draft_move(
            tenant=societe["tenant"],
            journal=societe["journal"],
            period=societe["periode"],
            date=dt.date(2026, 1, 10),
            narration="Achat mal ventilé",
        )
        add_line(
            brouillon,
            account=societe["charge"],
            label="Achat",
            debit=Decimal("500000"),
            analytic_distribution={"projet": {"INEXISTANT": 100}},
        )
        add_line(brouillon, account=societe["produit"], label="Vente", credit=Decimal("500000"))
        with pytest.raises(ValidationError) as refus:
            post_move(brouillon)
    assert "INEXISTANT" in str(refus.value)
    assert "projet" in str(refus.value)


# --------------------------------------------------------------------------
# Taux de change
# --------------------------------------------------------------------------


def test_un_taux_de_change_saisi_leve_le_refus_des_ecritures_en_devise(societe) -> None:
    tenant = societe["tenant"]
    with use_tenant(tenant.id), pytest.raises(ValidationError):
        get_rate(tenant, "EUR", dt.date(2026, 1, 15))

    reponse = societe["client"].post(
        "/accounting/config/exchange-rates/",
        {"currency": "eur", "date": "2026-01-10", "rate_to_mga": "4900,50", "source": "BCM"},
    )
    assert reponse.status_code == 200, reponse.content

    with use_tenant(tenant.id):
        taux = AccExchangeRate.objects.filter(currency="EUR").first()
        assert taux is not None, "Le taux saisi n'a pas été enregistré."
        assert taux.rate_to_mga == Decimal("4900.500000")
        assert get_rate(tenant, "EUR", dt.date(2026, 1, 15)) == Decimal("4900.500000")

    contenu = societe["client"].get("/accounting/config/exchange-rates/").content.decode()
    assert "EUR" in contenu
    assert "BCM" in contenu


def test_un_taux_illisible_est_refuse_sans_rien_ecrire(societe) -> None:
    reponse = societe["client"].post(
        "/accounting/config/exchange-rates/",
        {"currency": "USD", "date": "pas-une-date", "rate_to_mga": "4500"},
    )
    assert reponse.status_code == 200
    assert "Date ou taux illisible" in reponse.content.decode()
    with use_tenant(societe["tenant"].id):
        assert not AccExchangeRate.objects.filter(currency="USD").exists()


def test_les_ecrans_de_g4_sont_atteignables_depuis_leurs_hubs(societe) -> None:
    operations = societe["client"].get("/accounting/operations/").content.decode()
    for chemin in ("/accounting/landed-costs/", "/accounting/dcom/"):
        assert chemin in operations, f"Le hub d'opérations ne mène pas à {chemin}."

    configuration = societe["client"].get("/accounting/config/").content.decode()
    for chemin in (
        "/accounting/config/analytic-plans/",
        "/accounting/config/analytic-accounts/",
        "/accounting/config/exchange-rates/",
    ):
        assert chemin in configuration, f"Le hub de configuration ne mène pas à {chemin}."

    imports = societe["client"].get("/accounting/config/imports/").content.decode()
    assert "/accounting/config/imports/invoices/" in imports


def test_une_societe_sans_exercice_nouvre_pas_un_500_sur_la_dcom() -> None:
    """Le premier jour d'une société neuve, et le crawler l'a trouvé.

    `filter(fiscal_year_id="")` fait lever « n'est pas un UUID valide »,
    qu'aucun gestionnaire ne rattrape : la page rendait 500 tant qu'aucun
    exercice n'existait. Même famille que les identifiants non typés fermés
    en T4bis, sur la valeur VIDE cette fois — et jamais visible en
    relecture, parce que toutes les fixtures créent un exercice."""
    tenant = Tenant.objects.create(code="G4N", name="Societe neuve")
    user = User.objects.create_user(email="g4-neuve@example.com", password="Str0ngPassw0rd!23")
    grant_module_access(user, "accounting")
    UserTenantMembership.objects.get_or_create(
        user=user, tenant=tenant, defaults={"is_default": True}
    )
    client = Client()
    client.force_login(user)
    session = client.session
    session["tenant_id"] = str(tenant.id)
    session.save()

    reponse = client.get("/accounting/dcom/")
    assert reponse.status_code == 200
    assert "Aucune déclaration pour cet exercice" in reponse.content.decode()


# --------------------------------------------------------------------------
# G-4bis — les deux référentiels que la remesure a trouvés
# --------------------------------------------------------------------------


def test_une_regle_de_rapprochement_se_cree_et_se_relit(societe) -> None:
    """Le moteur de rapprochement lit ces règles à chaque suggestion, et
    personne ne pouvait les écrire : elles se semaient par import et ne se
    corrigeaient qu'en base."""
    from apps.accounting.models import AccReconcileRule

    with use_tenant(societe["tenant"].id):
        banque = AccAccount.objects.create(
            tenant=societe["tenant"],
            code="512000",
            name="Banque",
            account_class="5",
            type=AccAccount.TYPE_BANK,
        )

    reponse = societe["client"].post(
        "/accounting/config/reconcile-rules/",
        {
            "name": "Virements clients",
            "bank_account_id": str(banque.id),
            "match_on_amount": "1",
            "amount_tolerance_mga": "500",
            "match_on_reference": "1",
            "priority": "10",
        },
    )
    assert reponse.status_code == 200, reponse.content

    with use_tenant(societe["tenant"].id):
        regle = AccReconcileRule.objects.filter(name="Virements clients").first()
        assert regle is not None, "La règle n'a pas été écrite."
        assert regle.priority == 10
        assert regle.match_on_reference is True
        assert regle.amount_tolerance_mga == Decimal("500.0000")

    contenu = societe["client"].get("/accounting/config/reconcile-rules/").content.decode()
    assert "Virements clients" in contenu
    assert "512000" in contenu


def test_une_categorie_de_caisse_sassocie_une_seule_fois(societe) -> None:
    from apps.accounting.models import AccCashCategoryMapping

    donnees = {"category_label": "Vente au comptant", "account_id": str(societe["produit"].id)}
    premier = societe["client"].post("/accounting/config/cash-categories/", donnees)
    assert premier.status_code == 200, premier.content
    with use_tenant(societe["tenant"].id):
        assert (
            AccCashCategoryMapping.objects.filter(category_label="Vente au comptant").count() == 1
        )

    # La contrainte d'unicité porte sur (société, catégorie) : une seconde
    # association se refuse lisiblement plutôt que de doubler la règle.
    second = societe["client"].post("/accounting/config/cash-categories/", donnees)
    assert second.status_code == 200
    assert "déjà associée" in second.content.decode()
    with use_tenant(societe["tenant"].id):
        assert (
            AccCashCategoryMapping.objects.filter(category_label="Vente au comptant").count() == 1
        )

    contenu = societe["client"].get("/accounting/config/cash-categories/").content.decode()
    assert "Vente au comptant" in contenu
    assert "707000" in contenu
